from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook


EXAMPLE_DIR = Path("example")
EXPECTED_WORKBOOK = EXAMPLE_DIR / "vedomost_tables_test.xlsx"
DEFAULT_ACTUAL_DIR = EXAMPLE_DIR / "actual"
IGNORED_SHEETS = {"Metadata"}
DEFAULT_MIN_CELL_ACCURACY = 0.75
DEFAULT_MIN_SHEET_ACCURACY = 0.60
HEADER_MARKERS = ("№", "наименование", "ед.изм", "кол-во", "примечание")


@dataclass(frozen=True)
class SheetScore:
    name: str
    matched: int
    total: int
    accuracy: float
    fuzzy: float
    numeric_matched: int
    numeric_total: int
    missing: bool = False


def main() -> None:
    args = parse_args()
    expected = load_expected_tables(args.expected)
    actual = load_actual_tables(args.actual)
    scores = compare_tables(expected, actual)

    print_report(scores)

    overall_matched = sum(score.matched for score in scores)
    overall_total = sum(score.total for score in scores)
    overall_accuracy = overall_matched / overall_total if overall_total else 0
    worst_sheet_accuracy = min((score.accuracy for score in scores), default=0)

    if overall_accuracy < args.min_cell_accuracy:
        raise SystemExit(
            f"overall accuracy {overall_accuracy:.1%} is below required {args.min_cell_accuracy:.1%}"
        )

    if worst_sheet_accuracy < args.min_sheet_accuracy:
        raise SystemExit(
            f"worst sheet accuracy {worst_sheet_accuracy:.1%} is below required {args.min_sheet_accuracy:.1%}"
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Compare service XLSX output with expected example tables.")
    parser.add_argument("--expected", type=Path, default=EXPECTED_WORKBOOK)
    parser.add_argument("--actual", type=Path, default=DEFAULT_ACTUAL_DIR)
    parser.add_argument("--min-cell-accuracy", type=float, default=DEFAULT_MIN_CELL_ACCURACY)
    parser.add_argument("--min-sheet-accuracy", type=float, default=DEFAULT_MIN_SHEET_ACCURACY)
    return parser.parse_args()


def load_expected_tables(path: Path) -> dict[str, list[list[str]]]:
    if not path.exists():
        raise SystemExit(f"expected workbook not found: {path}")

    workbook = load_workbook(path, data_only=True)
    return {
        sheet_name: extract_comparable_table(read_rows(workbook[sheet_name]))
        for sheet_name in workbook.sheetnames
    }


def load_actual_tables(path: Path) -> dict[str, list[list[str]]]:
    if not path.exists():
        raise SystemExit(
            f"actual output not found: {path}. Export service workbooks into {DEFAULT_ACTUAL_DIR}/ first."
        )

    if path.is_file():
        return load_actual_workbook(path)

    actual: dict[str, list[list[str]]] = {}
    for workbook_path in sorted(path.glob("*.xlsx")):
        workbook_tables = load_actual_workbook(workbook_path)
        if len(workbook_tables) == 1:
            actual[workbook_path.stem] = next(iter(workbook_tables.values()))
            continue

        actual.update(workbook_tables)

    if not actual:
        raise SystemExit(f"no .xlsx files found in actual output directory: {path}")

    return actual


def load_actual_workbook(path: Path) -> dict[str, list[list[str]]]:
    workbook = load_workbook(path, data_only=True)
    tables: dict[str, list[list[str]]] = {}

    for sheet_name in workbook.sheetnames:
        if sheet_name in IGNORED_SHEETS:
            continue

        table = extract_comparable_table(read_rows(workbook[sheet_name]))
        if table:
            tables[sheet_name] = table

    if not tables:
        raise SystemExit(f"actual workbook has no comparable sheets: {path}")

    return tables


def read_rows(ws, start_row: int = 1) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in ws.iter_rows(min_row=start_row, values_only=True):
        rows.append([normalize_cell(value) for value in row])
    return rows


def trim_table(rows: list[list[str]]) -> list[list[str]]:
    while rows and not any(rows[-1]):
        rows.pop()

    if not rows:
        return []

    max_width = max((last_non_empty_index(row) for row in rows), default=0)
    return [row[:max_width] + [""] * max(0, max_width - len(row)) for row in rows]


def extract_comparable_table(rows: list[list[str]]) -> list[list[str]]:
    table = trim_table(rows)
    if not table:
        return []

    for index, row in enumerate(table):
        marker_count = sum(1 for marker in HEADER_MARKERS if any(marker in value for value in row))
        if marker_count >= 2:
            return trim_table(table[index:])

    return table


def last_non_empty_index(row: list[str]) -> int:
    for index in range(len(row) - 1, -1, -1):
        if row[index]:
            return index + 1
    return 0


def normalize_cell(value) -> str:
    if value is None:
        return ""

    text = str(value).replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text, flags=re.MULTILINE).strip().lower()
    return text


def compare_tables(expected: dict[str, list[list[str]]], actual: dict[str, list[list[str]]]) -> list[SheetScore]:
    scores: list[SheetScore] = []

    for sheet_name, expected_table in expected.items():
        actual_table = actual.get(sheet_name)
        if actual_table is None:
            total = count_cells(expected_table)
            scores.append(
                SheetScore(
                    name=sheet_name,
                    matched=0,
                    total=total,
                    accuracy=0,
                    fuzzy=0,
                    numeric_matched=0,
                    numeric_total=count_numeric_cells(expected_table),
                    missing=True,
                )
            )
            continue

        matched, total, fuzzy, numeric_matched, numeric_total = compare_table(expected_table, actual_table)
        accuracy = matched / total if total else 0
        scores.append(
            SheetScore(
                name=sheet_name,
                matched=matched,
                total=total,
                accuracy=accuracy,
                fuzzy=fuzzy,
                numeric_matched=numeric_matched,
                numeric_total=numeric_total,
            )
        )

    return scores


def compare_table(expected: list[list[str]], actual: list[list[str]]) -> tuple[int, int, float, int, int]:
    row_count = max(len(expected), len(actual))
    col_count = max(max_width(expected), max_width(actual))
    matched = 0
    total = 0
    fuzzy_sum = 0.0
    numeric_matched = 0
    numeric_total = 0

    for row_index in range(row_count):
        for col_index in range(col_count):
            expected_value = value_at(expected, row_index, col_index)
            actual_value = value_at(actual, row_index, col_index)

            if not expected_value and not actual_value:
                continue

            total += 1
            if expected_value == actual_value:
                matched += 1
            fuzzy_sum += cell_similarity(expected_value, actual_value)

            expected_number = normalize_number(expected_value)
            if expected_number:
                numeric_total += 1
                if expected_number == normalize_number(actual_value):
                    numeric_matched += 1

    return matched, total, fuzzy_sum / total if total else 0, numeric_matched, numeric_total


def count_cells(table: list[list[str]]) -> int:
    return sum(1 for row in table for value in row if value)


def count_numeric_cells(table: list[list[str]]) -> int:
    return sum(1 for row in table for value in row if normalize_number(value))


def max_width(table: list[list[str]]) -> int:
    return max((len(row) for row in table), default=0)


def value_at(table: list[list[str]], row: int, col: int) -> str:
    if row >= len(table) or col >= len(table[row]):
        return ""

    return table[row][col]


def cell_similarity(expected: str, actual: str) -> float:
    if expected == actual:
        return 1.0

    if not expected or not actual:
        return 0.0

    return SequenceMatcher(None, expected, actual).ratio()


def normalize_number(value: str) -> str:
    match = re.search(r"\d+(?:[,.]\d+)?", value)
    if not match:
        return ""

    return match.group(0).replace(".", ",")


def print_report(scores: list[SheetScore]) -> None:
    print("sheet,matched,total,accuracy,fuzzy,numeric,status")
    for score in scores:
        status = "missing" if score.missing else "ok"
        numeric = f"{score.numeric_matched}/{score.numeric_total}" if score.numeric_total else "0/0"
        print(f"{score.name},{score.matched},{score.total},{score.accuracy:.1%},{score.fuzzy:.1%},{numeric},{status}")

    matched = sum(score.matched for score in scores)
    total = sum(score.total for score in scores)
    accuracy = matched / total if total else 0
    fuzzy = weighted_average([(score.fuzzy, score.total) for score in scores])
    numeric_matched = sum(score.numeric_matched for score in scores)
    numeric_total = sum(score.numeric_total for score in scores)
    print(f"overall,{matched},{total},{accuracy:.1%},{fuzzy:.1%},{numeric_matched}/{numeric_total},summary")


def weighted_average(values: list[tuple[float, int]]) -> float:
    total_weight = sum(weight for _, weight in values)
    if not total_weight:
        return 0

    return sum(value * weight for value, weight in values) / total_weight


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit(130)
